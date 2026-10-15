"""Theme-aware DOCX renderer for the canonical RAS model."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from .components import component_colors, component_kind, component_variant, semantic_label
from .markdown import Element, parse_inline, parse_markdown
from .models import ReportProject


class RenderError(RuntimeError):
    """Raised when the optional DOCX runtime is unavailable."""


def render_docx(project: ReportProject, output_path: str | Path, visuals: dict[str, dict[str, str]] | None = None) -> Path:
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise RenderError("python-docx is required for DOCX output") from exc

    visuals = visuals or {}
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    design = project.design
    palette = design.get("palette", {}) or {}
    typography = design.get("typography", {}) or {}
    page = design.get("page", {}) or {}
    margins = page.get("margins", {}) or {}
    doc = Document()
    section = doc.sections[0]
    if project.style.page_geometry == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(float(margins.get("top", 0.62)) + 0.15)
    section.bottom_margin = Inches(float(margins.get("bottom", 0.62)) + 0.15)
    section.left_margin = Inches(float(margins.get("left", 0.68)))
    section.right_margin = Inches(float(margins.get("right", 0.68)))
    colors = palette
    fonts = {
        "body": typography.get("body", "Aptos"),
        "heading": typography.get("heading", "Aptos Display"),
        "mono": typography.get("mono", "Consolas"),
        "cjk_fallback": typography.get("cjk", "Microsoft JhengHei"),
    }
    _configure_styles(doc, fonts, colors, Pt, RGBColor, qn)
    _configure_header(section, project, design, fonts, colors, Pt, RGBColor, OxmlElement, qn)
    doc.core_properties.title = project.title
    doc.core_properties.subject = f"RAS {project.style.archetype}"
    doc.core_properties.author = "RAS — Report Anything Skill"

    _add_cover(doc, project, design, fonts, colors, Pt, RGBColor, WD_ALIGN_PARAGRAPH, WD_TABLE_ALIGNMENT, OxmlElement, qn)

    figure_index = 0
    for index, block in enumerate(project.blocks, start=1):
        if bool((block.spec.presentation or {}).get("break_before")) and index > 1:
            doc.add_page_break()
        heading = doc.add_heading(f"{index:02d}  {block.spec.title}", level=1)
        _keep_with_next(heading)
        kind = component_kind(block)
        if kind in {"summary", "finding", "recommendation"}:
            _add_semantic_panel(doc, block, design, colors, fonts, Pt, RGBColor, WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, OxmlElement, qn)
        elements = parse_markdown(block.body)
        if elements and elements[0].kind == "heading" and _same_heading(elements[0].text, block.spec.title):
            elements = elements[1:]
        figure = visuals.get(block.spec.id)
        visual_first = project.style.emphasis == "visual-first" or bool(block.spec.visual.get("position") == "before")
        if visual_first and figure:
            figure_index += 1
            _add_visual(doc, figure, block.spec.visual, figure_index, Inches, Pt, WD_ALIGN_PARAGRAPH)
        _add_elements(doc, elements, colors, fonts, Pt, RGBColor, WD_ALIGN_PARAGRAPH, WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, OxmlElement, qn)
        if not visual_first and figure:
            figure_index += 1
            _add_visual(doc, figure, block.spec.visual, figure_index, Inches, Pt, WD_ALIGN_PARAGRAPH)

    if project.sources:
        doc.add_heading("Sources", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, ("ID", "Source", "Details")):
            cell.text = value
            _shade_cell(cell, _hex(palette.get("accent_soft", "#EAF2FF")), OxmlElement, qn)
            _set_cell_text_color(cell, _hex(palette.get("ink", "#17202A")), RGBColor)
        for row_index, source in enumerate(project.sources, start=1):
            cells = table.add_row().cells
            cells[0].text = source.id
            cells[1].text = source.title
            cells[2].text = "; ".join(item for item in [source.author or source.publisher, source.url or source.file, source.publication_date] if item)
            if row_index % 2 == 0:
                for cell in cells:
                    _shade_cell(cell, _hex((project.brand.get("table", {}) or {}).get("stripe_fill", palette.get("surface", "#F5F7FA"))), OxmlElement, qn)

    assumptions = project.raw_manifest.get("assumptions", []) or []
    omitted = project.raw_manifest.get("omitted_material", []) or []
    if assumptions or omitted:
        doc.add_heading("Quality receipt", level=1)
        if assumptions:
            doc.add_paragraph("Assumptions", style="Heading 2")
            for item in assumptions:
                doc.add_paragraph(str(item), style="List Bullet")
        if omitted:
            doc.add_paragraph("Omitted or unresolved material", style="Heading 2")
            for item in omitted:
                doc.add_paragraph(str(item), style="List Bullet")

    _add_page_numbers(doc, project, design, qn, OxmlElement, Pt, RGBColor)
    doc.save(output)
    return output


def _add_cover(doc: Any, project: ReportProject, design: dict[str, Any], fonts: dict[str, Any], colors: dict[str, Any], Pt: Any, RGBColor: Any, WD_ALIGN_PARAGRAPH: Any, WD_TABLE_ALIGNMENT: Any, OxmlElement: Any, qn: Any) -> None:
    theme = str(design.get("theme_id", "minimal"))
    dark_cover = theme in {"institutional", "technical", "data-forward"}
    accent = _hex(colors.get("accent", "#2563EB"))
    paper = _hex(colors.get("paper", "#FFFFFF"))
    ink = _hex(paper if dark_cover else colors.get("ink", "#17202A"))
    muted = _hex(paper if dark_cover else colors.get("muted", "#5D6B78"))
    cover = doc.add_table(rows=4, cols=1)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.style = "Table Grid"
    fills = [accent if dark_cover else _hex(colors.get("surface", "#F5F7FA"))] * 4
    labels = [
        f"RAS / {design.get('display_name', 'REPORT').upper()}",
        project.title,
        f"{project.style.archetype.replace('-', ' ').title()}  ·  {project.style.tone.title()}  ·  {project.style.emphasis.replace('-', ' ').title()}",
        str((project.raw_manifest.get("style_decision", {}) or {}).get("style_receipt", project.style.rationale)),
    ]
    sizes = [8, 25, 10, 8]
    for index, (row, label, fill, size) in enumerate(zip(cover.rows, labels, fills, sizes)):
        cell = row.cells[0]
        cell.text = ""
        _shade_cell(cell, fill, OxmlElement, qn)
        _set_cell_border(cell, accent, OxmlElement, qn, left=8 if index == 0 else 1)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(8 if index == 0 else 2)
        paragraph.paragraph_format.space_after = Pt(11 if index in {1, 3} else 5)
        run = paragraph.add_run(label)
        run.font.name = fonts.get("heading" if index in {0, 1} else "body", "Aptos")
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(ink if index in {1, 2} else muted)
        if index == 1:
            run.bold = True
    doc.add_paragraph()


def _add_semantic_panel(doc: Any, block: Any, design: dict[str, Any], colors: dict[str, Any], fonts: dict[str, Any], Pt: Any, RGBColor: Any, WD_TABLE_ALIGNMENT: Any, WD_CELL_VERTICAL_ALIGNMENT: Any, OxmlElement: Any, qn: Any) -> None:
    kind = component_kind(block)
    accent, soft, _ = component_colors(design, kind)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    label_cell, body_cell = table.rows[0].cells
    label_cell.text = f"{semantic_label(kind)}\n{component_variant(block, design)}"
    body_cell.text = block.spec.goal or "Evidence-led report component"
    _shade_cell(label_cell, _hex(soft), OxmlElement, qn)
    _shade_cell(body_cell, _hex(soft), OxmlElement, qn)
    _set_cell_border(label_cell, _hex(accent), OxmlElement, qn, left=6)
    _set_cell_text_color(label_cell, _hex(accent), RGBColor)
    _set_cell_text_color(body_cell, _hex(colors.get("ink", "#17202A")), RGBColor)
    for cell in (label_cell, body_cell):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs:
                run.font.name = fonts.get("body", "Aptos")
                run.font.size = Pt(8.5 if cell is label_cell else 9)
    doc.add_paragraph()


def _configure_styles(doc: Any, fonts: dict[str, Any], colors: dict[str, Any], Pt: Any, RGBColor: Any, qn: Any) -> None:
    from docx.enum.style import WD_STYLE_TYPE

    body_font = str(fonts.get("body", "Aptos"))
    heading_font = str(fonts.get("heading", body_font))
    ink = RGBColor.from_string(_hex(colors.get("ink", "#17202A")))
    normal = doc.styles["Normal"]
    normal.font.name = body_font
    normal.font.size = Pt(10)
    normal.font.color.rgb = ink
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), str(fonts.get("cjk_fallback", "Microsoft JhengHei")))
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[name]
        style.font.name = heading_font
        style.font.color.rgb = RGBColor.from_string(_hex(colors.get("accent", "#2563EB")))
        style._element.rPr.rFonts.set(qn("w:eastAsia"), str(fonts.get("cjk_fallback", "Microsoft JhengHei")))
    if "RAS Caption" not in [style.name for style in doc.styles]:
        caption = doc.styles.add_style("RAS Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = body_font
        caption.font.size = Pt(8)
        caption.font.italic = True
        caption.font.color.rgb = RGBColor.from_string(_hex(colors.get("muted", "#5D6B78")))


def _configure_header(section: Any, project: ReportProject, design: dict[str, Any], fonts: dict[str, Any], colors: dict[str, Any], Pt: Any, RGBColor: Any, OxmlElement: Any, qn: Any) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = 0
    paragraph.text = f"{project.title}  ·  {design.get('display_name', 'RAS')}"
    for run in paragraph.runs:
        run.font.name = fonts.get("body", "Aptos")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(_hex(colors.get("muted", "#5D6B78")))
    _set_paragraph_border(paragraph, _hex(colors.get("accent", "#2563EB")), OxmlElement, qn)


def _add_elements(doc: Any, elements: list[Element], colors: dict[str, Any], fonts: dict[str, Any], Pt: Any, RGBColor: Any, WD_ALIGN_PARAGRAPH: Any, WD_TABLE_ALIGNMENT: Any, WD_CELL_VERTICAL_ALIGNMENT: Any, OxmlElement: Any, qn: Any) -> None:
    for element in elements:
        if element.kind == "heading":
            level = min(3, max(2, element.level))
            p = doc.add_paragraph(style=f"Heading {level}")
            _add_inline_runs(p, element.text, fonts, Pt)
            _keep_with_next(p)
        elif element.kind == "paragraph":
            p = doc.add_paragraph()
            _add_inline_runs(p, element.text, fonts, Pt)
        elif element.kind == "bullet_list":
            for item in element.items or []:
                p = doc.add_paragraph(style="List Bullet")
                _add_inline_runs(p, item, fonts, Pt)
        elif element.kind == "callout":
            p = doc.add_paragraph(style="Quote")
            _add_inline_runs(p, element.text, fonts, Pt, RGBColor, _hex(colors.get("accent", "#2563EB")), True)
        elif element.kind == "code":
            p = doc.add_paragraph()
            for run in p.runs:
                run.clear()
            run = p.add_run(element.text)
            run.font.name = str(fonts.get("mono", "Consolas"))
            run.font.size = Pt(8)
        elif element.kind == "table":
            rows = element.rows or []
            if not rows:
                continue
            columns = max(len(row) for row in rows)
            table = doc.add_table(rows=1, cols=columns)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for index in range(columns):
                cell = table.rows[0].cells[index]
                _set_cell_inline(cell, rows[0][index] if index < len(rows[0]) else "", fonts, Pt)
                _shade_cell(cell, _hex(colors.get("accent_soft", colors.get("surface", "#F4F6F8"))), OxmlElement, qn)
                _set_cell_text_color(cell, _hex(colors.get("ink", "#17202A")), RGBColor)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for row_index, row in enumerate(rows[1:], start=1):
                cells = table.add_row().cells
                for index in range(columns):
                    _set_cell_inline(cells[index], row[index] if index < len(row) else "", fonts, Pt)
                    cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if row_index % 2 == 0:
                        _shade_cell(cells[index], _hex(colors.get("surface", "#F8FAFC")), OxmlElement, qn)


def _add_inline_runs(paragraph: Any, text: str, fonts: dict[str, Any], Pt: Any, RGBColor: Any | None = None, color: str | None = None, bold_override: bool = False) -> None:
    for span in parse_inline(text):
        if not span.text:
            continue
        run = paragraph.add_run(span.text)
        run.bold = span.bold or bold_override
        run.italic = span.italic
        if span.code:
            run.font.name = str(fonts.get("mono", "Consolas"))
            run.font.size = Pt(8)
        if RGBColor is not None and color:
            run.font.color.rgb = RGBColor.from_string(_hex(color))


def _set_cell_inline(cell: Any, text: str, fonts: dict[str, Any], Pt: Any) -> None:
    cell.text = ""
    _add_inline_runs(cell.paragraphs[0], text, fonts, Pt)


def _add_visual(doc: Any, figure: dict[str, str], visual_spec: dict[str, Any], figure_index: int, Inches: Any, Pt: Any, WD_ALIGN_PARAGRAPH: Any) -> None:
    image_path = figure.get("png") or figure.get("svg")
    if not image_path or not Path(image_path).exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline = paragraph.add_run().add_picture(image_path, width=Inches(6.2))
    alt = str(visual_spec.get("alt_text") or figure.get("alt_text") or "Informative report visual")
    inline._inline.docPr.set("descr", alt)
    inline._inline.docPr.set("title", alt[:120])
    caption = str(visual_spec.get("caption", "")).strip()
    if caption:
        p = doc.add_paragraph(style="RAS Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"Figure {figure_index}. {caption}")
    alt_paragraph = doc.add_paragraph(style="RAS Caption")
    alt_paragraph.add_run(f"Alternative text: {alt}")


def _add_page_numbers(doc: Any, project: ReportProject, design: dict[str, Any], qn: Any, OxmlElement: Any, Pt: Any, RGBColor: Any) -> None:
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = 2
        run = footer.add_run(f"{design.get('theme_id', 'minimal')}  ·  Page ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        run._r.addnext(field)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(_hex((design.get("palette", {}) or {}).get("muted", "#5D6B78")))


def _shade_cell(cell: Any, fill: str, OxmlElement: Any, qn: Any) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill.replace("#", ""))
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_border(cell: Any, color: str, OxmlElement: Any, qn: Any, left: int = 1) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    edge = OxmlElement("w:left")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(left))
    edge.set(qn("w:color"), color.replace("#", ""))
    borders.append(edge)


def _set_cell_text_color(cell: Any, color: str, RGBColor: Any) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(_hex(color))


def _set_paragraph_border(paragraph: Any, color: str, OxmlElement: Any, qn: Any) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color.replace("#", ""))
    borders.append(bottom)
    p_pr.append(borders)


def _keep_with_next(paragraph: Any) -> None:
    paragraph.paragraph_format.keep_with_next = True


def _same_heading(left: str, right: str) -> bool:
    return re.sub(r"[^a-z0-9]", "", left.lower()) == re.sub(r"[^a-z0-9]", "", right.lower())


def _hex(value: str) -> str:
    text = str(value).replace("#", "")
    return (text + "FFFFFF")[:6]
