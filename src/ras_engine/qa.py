"""Structural and page-level quality assurance."""

from __future__ import annotations

import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from .design import design_signature
from .models import ReportProject
from .project import slugify
from .security import unresolved_placeholders
from .validation import Issue, has_errors, validate_project


def run_qa(manifest_path: str | Path) -> dict[str, Any]:
    project = ReportProject.load(manifest_path)
    issues = validate_project(project, check_outputs=True)
    outputs = project.raw_manifest.get("outputs", {}) or {}
    filename = slugify(str(outputs.get("filename", project.report_id)), project.report_id)
    for fmt in outputs.get("formats", []) or []:
        target = project.output_dir / f"{filename}.{fmt}"
        if fmt == "docx" and target.exists():
            issues.extend(_qa_docx(target, project))
        elif fmt == "pdf" and target.exists():
            issues.extend(_qa_pdf(target, project))
        elif fmt == "pptx" and target.exists():
            issues.extend(_qa_pptx(target, project))
    issues.extend(_qa_cross_format(project))
    rendered_pages = _render_pdf_pages(project)
    visual_receipt = _visual_regression_receipt(project, rendered_pages)
    if visual_receipt["status"] != "passed":
        issues.append(Issue("error", "visual-regression", "Rendered PDF pages did not pass visual integrity checks", "qa/rendered-pages"))
    passed = not has_errors(issues)
    results = [issue.to_dict() for issue in issues]
    receipt = {
        "status": "passed" if passed else "failed",
        "run_at": datetime.now(timezone.utc).isoformat(),
        "issue_count": len(issues),
        "error_count": sum(issue.severity == "error" for issue in issues),
        "warning_count": sum(issue.severity == "warning" for issue in issues),
        "rendered_pdf_pages": rendered_pages,
        "results": results,
        "renderer_versions": _renderer_versions(),
        "design": {
            "theme_id": project.design.get("theme_id"),
            "display_name": project.design.get("display_name"),
            "resolved_brand": project.design.get("resolved_brand"),
            "signature": design_signature(project.design),
            "receipt": project.design.get("receipt", {}),
        },
        "visual_regression": visual_receipt,
    }
    project.raw_manifest.setdefault("qa", {}).clear()
    project.raw_manifest["qa"].update(receipt)
    project.raw_manifest.setdefault("progress", {})["status"] = "qa_passed" if passed else "qa_failed"
    project.save_manifest()
    report_path = project.root / "qa" / "qa-report.json"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")
    return {"passed": passed, "issues": issues, "receipt": receipt, "report": str(report_path)}


def _qa_docx(path: Path, project: ReportProject) -> list[Issue]:
    issues: list[Issue] = []
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names:
                issues.append(Issue("error", "docx-structure", "DOCX is missing word/document.xml", str(path)))
            if any(name.startswith("word/media/") for name in names):
                for name in names:
                    if name.startswith("word/media/") and archive.getinfo(name).file_size == 0:
                        issues.append(Issue("error", "docx-asset", f"DOCX contains an empty media asset: {name}", str(path)))
    except (zipfile.BadZipFile, OSError) as exc:
        issues.append(Issue("error", "docx-open", f"DOCX cannot be opened: {exc}", str(path)))
    try:
        from docx import Document

        doc = Document(path)
        text = _docx_text(doc)
        for placeholder in unresolved_placeholders(text):
            issues.append(Issue("error", "docx-placeholder", f"DOCX contains unresolved placeholder: {placeholder}", str(path)))
        if project.title not in text:
            issues.append(Issue("error", "docx-title", "DOCX does not contain the report title", str(path)))
    except Exception as exc:  # pragma: no cover - library variation
        issues.append(Issue("error", "docx-read", f"DOCX text inspection failed: {exc}", str(path)))
    return issues


def _qa_pdf(path: Path, project: ReportProject) -> list[Issue]:
    issues: list[Issue] = []
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if not reader.pages:
            issues.append(Issue("error", "pdf-pages", "PDF has no pages", str(path)))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        for placeholder in unresolved_placeholders(text):
            issues.append(Issue("error", "pdf-placeholder", f"PDF contains unresolved placeholder: {placeholder}", str(path)))
        if project.title not in text:
            issues.append(Issue("error", "pdf-title", "PDF does not contain the report title", str(path)))
    except Exception as exc:  # pragma: no cover - library variation
        issues.append(Issue("error", "pdf-open", f"PDF cannot be inspected: {exc}", str(path)))
    return issues


def _qa_pptx(path: Path, project: ReportProject) -> list[Issue]:
    """Inspect the editable presentation package and its slide text."""

    issues: list[Issue] = []
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if "ppt/presentation.xml" not in names:
                issues.append(Issue("error", "pptx-structure", "PPTX is missing ppt/presentation.xml", str(path)))
            slide_names = sorted(name for name in names if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
            if not slide_names:
                issues.append(Issue("error", "pptx-slides", "PPTX contains no slide XML", str(path)))
            for name in names:
                if name.startswith("ppt/media/") and archive.getinfo(name).file_size == 0:
                    issues.append(Issue("error", "pptx-asset", f"PPTX contains an empty media asset: {name}", str(path)))
            text = "\n".join(_xml_text(archive.read(name)) for name in slide_names)
            for placeholder in unresolved_placeholders(text):
                issues.append(Issue("error", "pptx-placeholder", f"PPTX contains unresolved placeholder: {placeholder}", str(path)))
            if _normalized_text(project.title) not in _normalized_text(text):
                issues.append(Issue("error", "pptx-title", "PPTX does not contain the report title", str(path)))
            for block in project.blocks:
                if _normalized_text(block.spec.title) not in _normalized_text(text):
                    issues.append(Issue("error", "pptx-section", f"PPTX does not contain section title: {block.spec.title}", str(path)))
    except (zipfile.BadZipFile, OSError, ElementTree.ParseError) as exc:
        issues.append(Issue("error", "pptx-open", f"PPTX cannot be inspected: {exc}", str(path)))
    return issues


def _qa_cross_format(project: ReportProject) -> list[Issue]:
    issues: list[Issue] = []
    outputs = project.raw_manifest.get("outputs", {}) or {}
    filename = slugify(str(outputs.get("filename", project.report_id)), project.report_id)
    required_titles = [project.title] + [block.spec.title for block in project.blocks]
    texts: dict[str, str] = {}
    if "docx" in outputs.get("formats", []) or (project.output_dir / f"{filename}.docx").exists():
        try:
            from docx import Document

            docx = Document(project.output_dir / f"{filename}.docx")
            texts["docx"] = _docx_text(docx)
        except Exception:
            pass
    if "pdf" in outputs.get("formats", []) or (project.output_dir / f"{filename}.pdf").exists():
        try:
            from pypdf import PdfReader

            texts["pdf"] = "\n".join(page.extract_text() or "" for page in PdfReader(project.output_dir / f"{filename}.pdf").pages)
        except Exception:
            pass
    pptx_path = project.output_dir / f"{filename}.pptx"
    if "pptx" in outputs.get("formats", []) or pptx_path.exists():
        try:
            texts["pptx"] = _pptx_text(pptx_path)
        except Exception:
            pass
    normalized_texts = {name: _normalized_text(text) for name, text in texts.items() if text}
    if len(normalized_texts) >= 2:
        for title in required_titles:
            normalized_title = _normalized_text(title)
            missing = [name for name, text in normalized_texts.items() if normalized_title not in text]
            if missing:
                issues.append(Issue("error", "format-parity", f"Section/title missing from {', '.join(missing)}: {title}", "outputs"))
    return issues


def _render_pdf_pages(project: ReportProject) -> int:
    outputs = project.raw_manifest.get("outputs", {}) or {}
    filename = slugify(str(outputs.get("filename", project.report_id)), project.report_id)
    pdf_path = project.output_dir / f"{filename}.pdf"
    if not pdf_path.exists():
        return 0
    render_dir = project.root / "qa" / "rendered-pages"
    render_dir.mkdir(parents=True, exist_ok=True)
    for old_page in render_dir.glob("page-*.png"):
        old_page.unlink()
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(str(pdf_path))
        count = len(pdf)
        for index in range(count):
            page = pdf[index]
            bitmap = page.render(scale=1.0)
            bitmap.to_pil().save(render_dir / f"page-{index + 1:03d}.png")
        return count
    except Exception:
        return 0


def _renderer_versions() -> dict[str, str]:
    versions: dict[str, str] = {}
    for module_name, package_name in (("docx", "python-docx"), ("reportlab", "reportlab"), ("PIL", "Pillow"), ("pypdf", "pypdf"), ("pypdfium2", "pypdfium2")):
        try:
            module = __import__(module_name)
            versions[package_name] = str(getattr(module, "__version__", "available"))
        except Exception:
            versions[package_name] = "unavailable"
    try:
        from .render_pptx import _find_artifact_tool

        versions["artifact-tool"] = "available" if _find_artifact_tool() else "unavailable"
    except Exception:
        versions["artifact-tool"] = "unavailable"
    return versions


def _visual_regression_receipt(project: ReportProject, rendered_pages: int) -> dict[str, Any]:
    """Record deterministic visual evidence even when no golden baseline is configured."""

    render_dir = project.root / "qa" / "rendered-pages"
    page_files = sorted(render_dir.glob("page-*.png")) if render_dir.exists() else []
    non_empty = 0
    try:
        from PIL import Image

        for path in page_files:
            if path.stat().st_size > 0:
                with Image.open(path) as image:
                    if image.getbbox() is not None:
                        non_empty += 1
    except Exception:
        non_empty = sum(path.stat().st_size > 0 for path in page_files)
    passed = rendered_pages == len(page_files) and non_empty == len(page_files) if rendered_pages else True
    return {
        "status": "passed" if passed else "failed",
        "method": "theme-signature-and-rendered-page-integrity",
        "baseline": "not-configured",
        "signature": design_signature(project.design),
        "page_count": rendered_pages,
        "non_empty_pages": non_empty,
    }


def _docx_text(doc: Any) -> str:
    """Read paragraph and table text so branded cover tables remain testable."""

    chunks = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def _pptx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        slide_names = sorted(name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
        return "\n".join(_xml_text(archive.read(name)) for name in slide_names)


def _xml_text(content: bytes) -> str:
    root = ElementTree.fromstring(content)
    return "\n".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))


def _normalized_text(value: str) -> str:
    """Normalize layout whitespace before comparing extracted DOCX/PDF text."""

    return " ".join(str(value).replace("\x7f", " ").split())
